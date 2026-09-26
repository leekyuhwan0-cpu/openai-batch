"""
openai-batch/poll_batches.py

GitHub Actions에서 실행. 카드뉴스 번역 파이프라인
(bonus_model_interfaces.submit_translation_batch)이 제출해둔 OpenAI Batch API
작업 중 완료된 것을 찾아 결과를 각 채널 워크북의 TRANSLATE_DATA 탭에 채워넣는다.

트리거는 두 가지: (1) submit_translation_batch가 제출 직후 배치 완료를 직접
감지해 즉시 쏘는 workflow_dispatch(평소엔 이걸로 2~3분 내 처리), (2) 1시간
간격 schedule cron - (1)이 놓친 경우(GUI 종료 등)를 위한 저빈도 백업(2026-09-11).

- 대상 판별: status=="completed" AND metadata.contents_sheet_id가 있는 배치만
  (이 프로젝트가 제출한 배치는 전부 submit_translation_batch가 이 키를 채워서 만듦 -
  같은 OpenAI 계정의 다른 용도 배치와 섞이지 않게 하는 마커).
- 최근 3일 이내 생성된 배치만 훑는다(Batch API completion_window가 최대 24h라
  그보다 오래된 건 이미 completed거나 expired 상태로 끝났을 것 - 매 실행마다 전체
  이력을 다시 훑지 않기 위한 상한).
- 멱등성: TRANSLATE_DATA의 대상 셀이 이미 채워져 있으면 건너뛴다. 실제로 새로 쓴 값이
  하나도 없으면 시트 업로드 자체를 생략(불필요한 Drive 쓰기 방지).

Gemini QA 검증 레이어(2026-09-26 확정, project_번역검증_gemini_파이프라인_확정 참고):
GPT 배치 번역 결과를 TRANSLATE_DATA에 쓰기 전에 gemini-3.8-flash로 (category,
target_lang, field) 그룹 단위 배치검증 -> FAIL만 gpt-6-luna reasoning_effort=xhigh로
재번역 -> 재검증 -> 그래도 FAIL이면 gemini가 누적 실패사유를 참고해 직접 최종번역.
poll_batches.py는 별도 레포라 bonus_model_interfaces.py/bonus_sheet_io.py를 그대로
import할 수 없음(파일 자체가 없고, bonus_sheet_io.py는 top-level `import msvcrt`라
Linux 러너에서 죽음) - 그래서 TRANSLATE_PROMPTS/SOURCE_ACCOUNTS 조회도 이 파일 안의
download_workbook/ws_read_dicts를 그대로 재사용해 MAIN_SHEET_ID를 직접 읽는다(OAuth
자격증명은 bonus_sheet_io.py의 하드코딩 값과 동일해 기존 GOOGLE_* 시크릿으로 충분).
"""
import json
import os
import time
import uuid

import requests
from openpyxl import load_workbook
from openai import OpenAI

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
REQUEST_TIMEOUT = 30
HERE = os.path.dirname(os.path.abspath(__file__))
MAX_BATCH_AGE_SEC = 3 * 24 * 3600
MAIN_SHEET_ID = "1IK7MbiTFJiV_ofeco9FwYs7UFY_f4QHHIDZJb7T_4-o"

GEMINI_MODEL = "gemini-3.8-flash"
GEMINI_KEYS = [os.environ[k] for k in (f"GEMINI_API_KEY_{i}" for i in range(1, 10)) if os.environ.get(k)]
XHIGH_MODEL = "gpt-6-luna"  # project_번역모델_확정 참고 - 후킹/캡션 둘다 이 모델로 확정운영중

GOOGLE_CLIENT_ID = os.environ["GOOGLE_CLIENT_ID"]
GOOGLE_CLIENT_SECRET = os.environ["GOOGLE_CLIENT_SECRET"]
GOOGLE_REFRESH_TOKEN = os.environ["GOOGLE_REFRESH_TOKEN"]

_access_token = None


def _get_access_token(force_refresh=False):
    global _access_token
    if _access_token and not force_refresh:
        return _access_token
    r = requests.post("https://oauth2.googleapis.com/token", data={
        "client_id": GOOGLE_CLIENT_ID,
        "client_secret": GOOGLE_CLIENT_SECRET,
        "refresh_token": GOOGLE_REFRESH_TOKEN,
        "grant_type": "refresh_token",
    }, timeout=REQUEST_TIMEOUT)
    r.raise_for_status()
    _access_token = r.json()["access_token"]
    return _access_token


def _request_with_retry(method, url, **kwargs):
    kwargs.setdefault("timeout", REQUEST_TIMEOUT)
    headers = kwargs.pop("headers", {})
    headers["Authorization"] = f"Bearer {_get_access_token()}"

    for attempt in range(1, 4):
        try:
            resp = requests.request(method, url, headers=headers, **kwargs)
        except (requests.exceptions.ConnectionError, requests.exceptions.SSLError, requests.exceptions.Timeout):
            if attempt == 3:
                raise
            time.sleep(2 * attempt)
            continue
        if resp.status_code == 401:
            headers["Authorization"] = f"Bearer {_get_access_token(force_refresh=True)}"
            continue
        if resp.status_code in (502, 503, 504) and attempt < 3:
            time.sleep(2 * attempt)
            continue
        resp.raise_for_status()
        return resp
    resp.raise_for_status()
    return resp


def download_workbook(file_id):
    r = _request_with_retry(
        "GET", f"https://www.googleapis.com/drive/v3/files/{file_id}/export",
        params={"mimeType": XLSX_MIME},
    )
    xlsx_path = os.path.join(HERE, f"_poll_io_{file_id}_{uuid.uuid4().hex}.xlsx")
    with open(xlsx_path, "wb") as f:
        f.write(r.content)
    return load_workbook(xlsx_path), xlsx_path


def upload_workbook(wb, xlsx_path, file_id):
    wb.save(xlsx_path)
    try:
        with open(xlsx_path, "rb") as f:
            content = f.read()
        _request_with_retry(
            "PATCH", f"https://www.googleapis.com/upload/drive/v3/files/{file_id}?uploadType=media",
            headers={"Content-Type": XLSX_MIME}, data=content,
        )
    finally:
        os.remove(xlsx_path)


def ws_read_dicts(ws):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = list(rows[0])
    out = []
    for row in rows[1:]:
        if all(v is None for v in row):
            continue
        out.append({headers[i]: row[i] for i in range(len(headers)) if i < len(row)})
    return headers, out


def ws_write_dicts(ws, headers, dict_rows):
    if ws.max_row >= 2:
        ws.delete_rows(2, ws.max_row - 1)
    for row in dict_rows:
        ws.append([row.get(h, "") for h in headers])


def upsert_translate_data(wb, xlsx_path, contents_sheet_id, updates):
    """updates: [(file_id, column, value), ...]. pipeline/bonus_ui.py의
    persist_fields_batch와 동일한 upsert 규칙(행이 없으면 새로 만든다) +
    이미 채워진 셀은 덮어쓰지 않는 멱등성 체크. wb/xlsx_path는 호출부가 CONTENTS/ASSETS
    조회에 이미 쓴 것과 같은 다운로드를 그대로 넘겨받는다(같은 워크북 중복 다운로드
    방지, 2026-09-26 Gemini 검증 레이어 추가하며 리팩터링). Returns: written_count."""
    ws = wb["TRANSLATE_DATA"]
    headers, rows = ws_read_dicts(ws)
    by_file_id = {r.get("file_id"): r for r in rows}
    written = 0
    for file_id, column, value in updates:
        r = by_file_id.get(file_id)
        if r is None:
            r = {h: "" for h in headers}
            r["file_id"] = file_id
            rows.append(r)
            by_file_id[file_id] = r
        if (r.get(column) or "").strip():
            continue  # 이미 채워져 있음 - 덮어쓰지 않음(멱등성)
        r[column] = value
        written += 1
    if written:
        ws_write_dicts(ws, headers, rows)
        upload_workbook(wb, xlsx_path, contents_sheet_id)
    else:
        os.remove(xlsx_path)
    return written


_main_lookup_cache = None


def _load_main_lookups():
    """MAIN_SHEET_ID(TRANSLATE_PROMPTS/SOURCE_ACCOUNTS가 있는 마스터 시트)를 실행당
    1회만 다운로드해 캐싱. category는 CONTENTS/ASSETS가 아니라 SOURCE_ACCOUNTS 행에
    있고(bonus_job_processor.py의 ch_data["sa"].get("category") 확인), SOURCE_ACCOUNTS는
    channel_id 단위지 contents_sheet_id 단위가 아니므로 contents_sheet_id ->
    category 매핑을 SOURCE_ACCOUNTS.contents_sheet_id 컬럼으로 만들어둔다."""
    global _main_lookup_cache
    if _main_lookup_cache is None:
        wb, xlsx_path = download_workbook(MAIN_SHEET_ID)
        try:
            _, tp_rows = ws_read_dicts(wb["TRANSLATE_PROMPTS"])
            _, sa_rows = ws_read_dicts(wb["SOURCE_ACCOUNTS"])
        finally:
            os.remove(xlsx_path)
        category_by_sheet_id = {r.get("contents_sheet_id"): r.get("category") for r in sa_rows}
        _main_lookup_cache = (tp_rows, category_by_sheet_id)
    return _main_lookup_cache


def _get_translate_prompt_row(tp_rows, category, lang):
    for row in tp_rows:
        if row.get("translate_category") == category and row.get("translate_lang") == lang:
            return row
    return None


def build_source_lookup(wb):
    """같은 contents_sheet_id 워크북 안의 CONTENTS/ASSETS 탭에서 custom_id
    (=ASSETS.local_file_id, bonus_ui.py의 submit_translation_batch 호출부 확인)로
    원문을 조회할 수 있는 딕셔너리를 만든다."""
    _, c_rows = ws_read_dicts(wb["CONTENTS"])
    _, a_rows = ws_read_dicts(wb["ASSETS"])
    content_by_id = {c.get("content_id"): c for c in c_rows}
    assets_by_local_file_id = {a.get("local_file_id"): a for a in a_rows if a.get("local_file_id")}
    return content_by_id, assets_by_local_file_id


def get_source_text(field, custom_id, content_by_id, assets_by_local_file_id):
    """field="hook"이면 그 asset 자신의 source_hook, "caption"이면 그 asset의
    content_id로 CONTENTS를 찾아 source_caption(캡션 배치의 custom_id는 main_asset의
    local_file_id라 이렇게 한 단계 거쳐야 함)."""
    asset = assets_by_local_file_id.get(custom_id)
    if asset is None:
        return None
    if field == "hook":
        return (asset.get("source_hook") or "").strip()
    content = content_by_id.get(asset.get("content_id"))
    return (content.get("source_caption") or "").strip() if content else None


def _gemini_call(payload):
    """9키 로테이션: 429(한도소진)면 다음 키로, 500/503(일시장애)이면 같은 키로 백오프
    재시도. 전부 실패하면 None(호출부가 검증 없이 원본 채택하도록)."""
    if not GEMINI_KEYS:
        print("  경고: GEMINI_API_KEY_1~9 중 설정된 키가 없음 - 검증 없이 원본 채택")
        return None
    for key_idx, api_key in enumerate(GEMINI_KEYS):
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={api_key}"
        for attempt in range(1, 4):
            try:
                resp = requests.post(url, json=payload, timeout=180)
            except (requests.exceptions.ConnectionError, requests.exceptions.Timeout):
                time.sleep(5)
                continue
            if resp.status_code == 200:
                return resp.json()
            if resp.status_code == 429:
                print(f"  Gemini key#{key_idx + 1} 429(한도소진) - 다음 키로 전환")
                break
            if resp.status_code in (500, 503):
                time.sleep(5 * attempt)
                continue
            print(f"  Gemini key#{key_idx + 1} HTTP {resp.status_code}: {resp.text[:200]}")
            break
    print("  경고: Gemini 호출 9키 전부 실패 - 검증 없이 원본 채택")
    return None


_VERIFY_SCHEMA = {
    "type": "ARRAY",
    "items": {
        "type": "OBJECT",
        "properties": {
            "id": {"type": "INTEGER"},
            "meaning": {"type": "STRING", "enum": ["accurate", "partial_loss", "distorted", "added"]},
            "numbers_names": {"type": "STRING", "enum": ["all_correct", "partial_error", "error"]},
            "subject_object": {"type": "STRING", "enum": ["same", "uncertain", "different"]},
            "verdict": {"type": "STRING", "enum": ["PASS", "FAIL"]},
            "reason": {"type": "STRING"},
        },
        "required": ["id", "meaning", "numbers_names", "subject_object", "verdict", "reason"],
    },
}


def _gemini_verify(items, style_hint, exchange_rate, field):
    """items: [{"id": int, "source": str, "translated": str}, ...], 같은
    (category, target_lang, field) 그룹 전체를 한 번에 검증(gemini_test6_ctx.py로
    100건 규모까지 검증된 패턴 그대로 재사용). Returns: {id: {verdict, reason, ...}}
    또는 실패 시 None."""
    field_label = "짧은 후킹 문구(제목형)" if field == "hook" else "본문(캡션)"
    rule_text = style_hint.replace("{EXCHANGE_RATE}", exchange_rate).replace(
        "{TEXT}", "(실제 번역대상 텍스트는 아래 검증항목에 개별 포함됨)"
    )
    context_text = f"### 번역 시 실제로 GPT에게 주어진 지침 ###\n{rule_text}"
    lines = [f"[id={it['id']}]\n원문: {it['source']}\n번역: {it['translated']}" for it in items]
    body_text = "\n\n".join(lines)
    prompt = f"""아래는 번역 시 실제로 사용된 스타일/화폐표기 지침이다. 이 지침에 맞게 작성된 번역은 정상으로 간주하고, 이 지침을 벗어난 경우만 오류로 판단하라.

{context_text}

===

위 지침을 참고해서, 다음 (원문, 번역) 쌍 {len(items)}개의 번역 품질을 채점하라. 원문은 인스타그램 카드뉴스용 {field_label}이다.

- meaning: accurate / partial_loss / distorted / added
- numbers_names: all_correct / partial_error / error (단, 위 지침에 명시된 화폐 축약/반올림 표기는 오류로 치지 않는다)
- subject_object: same / uncertain / different
- verdict: PASS / FAIL (위 3개 중 하나라도 문제있으면 FAIL)
- reason: FAIL일 때만 간단히 한 문장 (PASS면 빈 문자열)

{body_text}
"""
    payload = {
        "contents": [{"parts": [{"text": prompt}]}],
        "generationConfig": {"responseMimeType": "application/json", "responseSchema": _VERIFY_SCHEMA},
    }
    data = _gemini_call(payload)
    if data is None:
        return None
    try:
        text = data["candidates"][0]["content"]["parts"][0]["text"]
        return {p["id"]: p for p in json.loads(text)}
    except (KeyError, IndexError, TypeError, json.JSONDecodeError) as e:
        print(f"  경고: Gemini 검증 응답 파싱 실패({e}) - 검증 없이 원본 채택")
        return None


def _gemini_translate_fallback(source_text, attempts, reasons, style_hint, exchange_rate, field, target_lang):
    """xhigh 재번역까지 FAIL난 항목의 최종 수단: gemini가 누적 실패사유를 컨텍스트로
    받아 직접 최종번역까지 수행. Returns: str 또는 실패 시 None(호출부가 xhigh 재번역
    결과를 그대로 채택하도록)."""
    field_label = "짧은 후킹 문구(제목형)" if field == "hook" else "본문(캡션)"
    rule_text = style_hint.replace("{EXCHANGE_RATE}", exchange_rate).replace("{TEXT}", source_text)
    history = "\n".join(
        f"- 시도 {i + 1}: \"{t}\" -> 검증 실패사유: {r}" for i, (t, r) in enumerate(zip(attempts, reasons))
    )
    prompt = f"""다음은 인스타그램 카드뉴스용 {field_label} 번역 지침이다. 이 지침에 따라 아래 원문을 target_lang={target_lang}로 직접 번역하라.

{rule_text}

===

원문: {source_text}

앞서 GPT가 이 원문을 번역 시도했으나 모두 검증에서 실패했다. 실패 이력:
{history}

위 실패 원인들을 반드시 피해서, 정확하고 자연스러운 최종 번역문 하나만 출력하라. 다른 설명 없이 번역문 텍스트만 출력할 것."""
    data = _gemini_call({"contents": [{"parts": [{"text": prompt}]}]})
    if data is None:
        return None
    try:
        return data["candidates"][0]["content"]["parts"][0]["text"].strip()
    except (KeyError, IndexError, TypeError):
        return None


def _gpt_retry_xhigh(client, style_hint, exchange_rate, source_text):
    prompt = style_hint.replace("{EXCHANGE_RATE}", exchange_rate).replace("{TEXT}", source_text)
    resp = client.chat.completions.create(
        model=XHIGH_MODEL,
        messages=[{"role": "user", "content": prompt}],
        reasoning_effort="xhigh",
    )
    return (resp.choices[0].message.content or "").strip()


def verify_and_correct(client, contents_sheet_id, wb, raw_updates):
    """raw_updates: [(custom_id, field, target_lang, text), ...] (GPT 배치 원본 결과).
    (category, target_lang, field) 그룹별로 gemini-3.8-flash 검증 -> FAIL만 xhigh
    재번역 -> 재검증 -> 그래도 FAIL이면 gemini 최종번역. Returns: [(custom_id, column,
    value), ...] (upsert_translate_data에 바로 넘길 최종 형태)."""
    content_by_id, assets_by_local_file_id = build_source_lookup(wb)
    tp_rows, category_by_sheet_id = _load_main_lookups()
    category = category_by_sheet_id.get(contents_sheet_id)

    groups = {}
    for custom_id, field, target_lang, text in raw_updates:
        groups.setdefault((target_lang, field), []).append((custom_id, text))

    final_updates = []
    for (target_lang, field), pairs in groups.items():
        column = f"{target_lang}_{field}"
        tp_row = _get_translate_prompt_row(tp_rows, category, target_lang)
        style_hint = ""
        exchange_rate = ""
        if tp_row:
            style_hint = ((tp_row.get("style_hint") if field == "hook" else tp_row.get("caption_style_hint")) or "").strip()
            exchange_rate = (tp_row.get("exchange_rate") or "").strip()

        verify_items = []
        for idx, (custom_id, text) in enumerate(pairs):
            source_text = get_source_text(field, custom_id, content_by_id, assets_by_local_file_id)
            if not source_text:
                print(f"  경고: custom_id={custom_id} 원문 조회 실패 - 검증 없이 원본 채택")
                final_updates.append((custom_id, column, text))
                continue
            verify_items.append({"id": idx, "custom_id": custom_id, "source": source_text, "translated": text})

        if not verify_items:
            continue
        if not style_hint:
            print(f"  경고: category={category!r} lang={target_lang!r} field={field!r} style_hint 없음 - 검증 스킵")
            final_updates.extend((it["custom_id"], column, it["translated"]) for it in verify_items)
            continue

        results = _gemini_verify(
            [{"id": it["id"], "source": it["source"], "translated": it["translated"]} for it in verify_items],
            style_hint, exchange_rate, field,
        )
        if results is None:
            final_updates.extend((it["custom_id"], column, it["translated"]) for it in verify_items)
            continue

        fail_items = []
        for it in verify_items:
            r = results.get(it["id"])
            if r is None or r.get("verdict") == "PASS":
                final_updates.append((it["custom_id"], column, it["translated"]))
            else:
                fail_items.append(dict(it, reason=r.get("reason", "")))

        if not fail_items:
            print(f"  category={category} lang={target_lang} field={field}: {len(verify_items)}건 전부 PASS")
            continue
        print(f"  category={category} lang={target_lang} field={field}: {len(fail_items)}건 FAIL - xhigh 재번역")

        retry_items = []
        for it in fail_items:
            try:
                retried = _gpt_retry_xhigh(client, style_hint, exchange_rate, it["source"])
            except Exception as e:
                print(f"    xhigh 재번역 실패(custom_id={it['custom_id']}): {e}")
                final_updates.append((it["custom_id"], column, it["translated"]))
                continue
            retry_items.append(dict(it, retried=retried))

        if not retry_items:
            continue

        reverify_results = _gemini_verify(
            [{"id": it["id"], "source": it["source"], "translated": it["retried"]} for it in retry_items],
            style_hint, exchange_rate, field,
        )
        for it in retry_items:
            r = reverify_results.get(it["id"]) if reverify_results else None
            if reverify_results is None or r is None or r.get("verdict") == "PASS":
                final_updates.append((it["custom_id"], column, it["retried"]))
                continue
            print(f"    custom_id={it['custom_id']}: xhigh 재번역도 FAIL - gemini 최종번역 시도")
            fallback = _gemini_translate_fallback(
                it["source"], [it["translated"], it["retried"]], [it["reason"], r.get("reason", "")],
                style_hint, exchange_rate, field, target_lang,
            )
            final_updates.append((it["custom_id"], column, fallback or it["retried"]))

    return final_updates


def main():
    client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])
    now = time.time()

    batches = [
        b for b in client.batches.list(limit=100)
        if now - b.created_at <= MAX_BATCH_AGE_SEC
    ]
    completed = [
        b for b in batches
        if b.status == "completed" and (b.metadata or {}).get("contents_sheet_id") and b.output_file_id
    ]
    print(f"최근 {MAX_BATCH_AGE_SEC // 3600}h 배치 {len(batches)}개 중 처리대상 {len(completed)}개")

    by_sheet = {}
    for b in completed:
        by_sheet.setdefault(b.metadata["contents_sheet_id"], []).append(b)

    for contents_sheet_id, sheet_batches in by_sheet.items():
        raw_updates = []  # (custom_id, field, target_lang, text) - 검증 전 GPT 원본
        for b in sheet_batches:
            field = b.metadata.get("field")
            target_lang = b.metadata.get("target_lang")
            if not field or not target_lang:
                continue
            content = client.files.content(b.output_file_id).text
            for line in content.splitlines():
                if not line.strip():
                    continue
                row = json.loads(line)
                custom_id = row.get("custom_id")
                resp = row.get("response") or {}
                if resp.get("status_code") != 200:
                    print(f"  경고: batch={b.id} custom_id={custom_id} 응답 실패, 건너뜀")
                    continue
                try:
                    text = resp["body"]["choices"][0]["message"]["content"].strip()
                except (KeyError, IndexError, TypeError):
                    continue
                if custom_id and text:
                    raw_updates.append((custom_id, field, target_lang, text))

        if not raw_updates:
            continue

        wb, xlsx_path = download_workbook(contents_sheet_id)
        final_updates = verify_and_correct(client, contents_sheet_id, wb, raw_updates)
        if final_updates:
            written = upsert_translate_data(wb, xlsx_path, contents_sheet_id, final_updates)
            print(f"{contents_sheet_id}: {len(final_updates)}건 확인, {written}건 신규 반영")
        else:
            os.remove(xlsx_path)


if __name__ == "__main__":
    main()
